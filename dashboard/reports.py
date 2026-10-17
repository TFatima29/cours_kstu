"""Отчёт по студентам: страница с таблицей и выгрузка в Excel и PDF."""
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from django.contrib.auth import get_user_model
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from courses.models import Course
from learning.services import build_course_state
from quizzes.models import QuizAttempt

from .views import staff_required

User = get_user_model()

COLUMNS = ["№", "ФИО", "Логин", "E-mail", "Статус", "Курсы", "Прогресс, %", "Пройдено курсов", "Сдано тестов", "Средний балл, %", "Добавлен"]

# Шрифты с кириллицей: Windows, затем Linux (в Docker ставится fonts-dejavu-core)
FONT_CANDIDATES = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/Library/Fonts/Arial Unicode.ttf", "/Library/Fonts/Arial Unicode.ttf"),
]


def _selected_course(request):
    pk = request.GET.get("course", "")
    return Course.objects.filter(pk=pk).first() if pk.isdigit() else None


def build_rows(course=None):
    """Строка отчёта на каждого студента. Если выбран курс — только его слушатели и прогресс по нему."""
    students = User.objects.filter(is_staff=False).order_by("full_name")
    if course:
        students = students.filter(enrollments__course=course)
    rows = []
    for i, user in enumerate(students.prefetch_related("enrollments__course"), start=1):
        enrollments = [e for e in user.enrollments.all() if course is None or e.course_id == course.pk]
        states = [build_course_state(user, e.course, e) for e in enrollments]
        attempts = QuizAttempt.objects.filter(user=user)
        if course:
            attempts = attempts.filter(level__course=course)
        best = {}
        for level_id, score in attempts.values_list("level_id", "score"):
            best[level_id] = max(score, best.get(level_id, 0))
        rows.append(
            [
                i,
                user.full_name,
                user.login or "",
                user.email,
                "активен" if user.is_active else "заблокирован",
                ", ".join(e.course.title for e in enrollments) or "—",
                round(sum(s.percent for s in states) / len(states)) if states else 0,
                sum(1 for e in enrollments if e.is_completed),
                attempts.filter(passed=True).values("level").distinct().count(),
                round(sum(best.values()) / len(best)) if best else 0,
                timezone.localtime(user.date_joined).strftime("%d.%m.%Y"),
            ]
        )
    return rows


def _title(course):
    return f"Отчёт по студентам — {course.title}" if course else "Отчёт по студентам"


def _filename(course, ext):
    stamp = timezone.localtime().strftime("%Y-%m-%d")
    return f"students_{course.pk if course else 'all'}_{stamp}.{ext}"


@staff_required
def report_page(request):
    course = _selected_course(request)
    return render(
        request,
        "dashboard/reports.html",
        {"rows": build_rows(course), "columns": COLUMNS, "courses": Course.objects.order_by("title"), "course": course},
    )


@staff_required
def report_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    course = _selected_course(request)
    rows = build_rows(course)
    wb = Workbook()
    ws = wb.active
    ws.title = "Студенты"
    ws.append([_title(course)])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Сформирован {timezone.localtime().strftime('%d.%m.%Y %H:%M')}, студентов: {len(rows)}"])
    ws.append([])
    ws.append(COLUMNS)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    widths = [5, 28, 16, 28, 14, 40, 12, 12, 12, 14, 12]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=header_row + 1):
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(course, "xlsx")}"'
    return response


def _register_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "CP-Regular" in pdfmetrics.getRegisteredFontNames():
        return "CP-Regular", "CP-Bold"
    for regular, bold in FONT_CANDIDATES:
        if Path(regular).exists():
            pdfmetrics.registerFont(TTFont("CP-Regular", regular))
            pdfmetrics.registerFont(TTFont("CP-Bold", bold if Path(bold).exists() else regular))
            return "CP-Regular", "CP-Bold"
    return "Helvetica", "Helvetica-Bold"  # без кириллического шрифта русские буквы не отобразятся


@staff_required
def report_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    course = _selected_course(request)
    rows = build_rows(course)
    regular, bold = _register_font()
    cell = ParagraphStyle("cell", fontName=regular, fontSize=8, leading=10)
    head = ParagraphStyle("head", parent=cell, fontName=bold, textColor=colors.white)
    title = ParagraphStyle("title", fontName=bold, fontSize=15, leading=19)
    sub = ParagraphStyle("sub", fontName=regular, fontSize=9, textColor=colors.grey)

    data = [[Paragraph(c, head) for c in COLUMNS]]
    data += [[Paragraph(escape(str(v)), cell) for v in row] for row in rows]
    widths = [8, 40, 24, 42, 20, 62, 16, 16, 15, 17, 18]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F5597")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B4BCCB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F4F9")]),
            ]
        )
    )
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title=_title(course),
    )
    story = [
        Paragraph(escape(_title(course)), title),
        Paragraph(f"Course Platform. Сформирован {timezone.localtime().strftime('%d.%m.%Y %H:%M')}, студентов: {len(rows)}", sub),
        Spacer(1, 6 * mm),
        table if rows else Paragraph("Студентов нет.", cell),
    ]
    doc.build(story)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_filename(course, "pdf")}"'
    return response
