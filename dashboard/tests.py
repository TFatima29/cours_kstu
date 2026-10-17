"""Студенты, доступ к курсам, материалы, отчёты, вход по логину."""
import re

from django.contrib.auth import get_user_model
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from courses.models import Course, LibraryItem
from learning.models import Enrollment
from learning.tests import make_course

User = get_user_model()


@override_settings(MEDIA_ROOT="test_media")
class StudentsAdminTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user("adm@b.kg", "Strong-pass-42", full_name="Админ", is_staff=True)
        self.client.force_login(self.admin)
        self.course = make_course(levels=1, lessons=1, questions=1)
        self.other = Course.objects.create(title="Другой курс", is_published=True, price=2500)

    def test_create_student_with_login_and_course_access(self):
        r = self.client.post(
            reverse("dashboard:user_create"),
            {"full_name": "Нургуль", "email": "nurgul.a@mail.kg", "courses": [self.course.pk]},
        )
        user = User.objects.get(email="nurgul.a@mail.kg")
        self.assertRedirects(r, reverse("dashboard:user_detail", args=[user.pk]), fetch_redirect_response=False)
        self.assertEqual(user.login, "nurgul.a")
        self.assertEqual(list(user.enrollments.values_list("course_id", flat=True)), [self.course.pk])
        body = mail.outbox[0].body
        self.assertIn("Логин: nurgul.a", body)
        self.assertIn("Тестовый курс", body)
        password = re.search(r"Пароль: (\S+)", body).group(1)
        # студент входит по логину и видит только выданный курс
        self.client.logout()
        r = self.client.post(reverse("accounts:login"), {"username": "nurgul.a", "password": password})
        self.assertRedirects(r, reverse("learning:my_courses"), fetch_redirect_response=False)
        r = self.client.get(reverse("courses:catalog"))
        self.assertContains(r, "Тестовый курс")
        self.assertNotContains(r, "Другой курс")
        self.assertRedirects(self.client.get(self.other.get_absolute_url()), reverse("learning:my_courses"))
        # самостоятельной записи нет
        self.client.post(reverse("courses:enroll", args=[self.other.pk]))
        self.assertFalse(Enrollment.objects.filter(user=user, course=self.other).exists())

    def test_generated_login_is_unique(self):
        User.objects.create_user("x@a.kg", "Strong-pass-42", full_name="X", login="ivan")
        self.client.post(reverse("dashboard:user_create"), {"full_name": "Иван", "email": "ivan@b.kg"})
        login = User.objects.get(email="ivan@b.kg").login
        self.assertTrue(login.startswith("ivan") and login != "ivan")

    def test_edit_user_login_email_and_password(self):
        user = User.objects.create_user("u@b.kg", "Old-pass-42", full_name="U", login="u_old")
        r = self.client.post(
            reverse("dashboard:user_edit", args=[user.pk]),
            {"full_name": "Новое Имя", "email": "new@b.kg", "login": "u_new", "is_active": "on",
             "new_password": "Brand-new-99", "send_email": "on"},
        )
        self.assertRedirects(r, reverse("dashboard:user_detail", args=[user.pk]), fetch_redirect_response=False)
        user.refresh_from_db()
        self.assertEqual((user.full_name, user.email, user.login), ("Новое Имя", "new@b.kg", "u_new"))
        self.assertTrue(user.check_password("Brand-new-99"))
        self.assertEqual(mail.outbox[0].to, ["new@b.kg"])

    def test_edit_rejects_taken_login(self):
        User.objects.create_user("a@b.kg", "Strong-pass-42", full_name="A", login="busy")
        user = User.objects.create_user("u@b.kg", "Strong-pass-42", full_name="U")
        r = self.client.post(reverse("dashboard:user_edit", args=[user.pk]), {"full_name": "U", "email": "u@b.kg", "login": "BUSY"})
        self.assertContains(r, "уже занят")

    def test_deactivate_and_activate(self):
        user = User.objects.create_user("u@b.kg", "Strong-pass-42", full_name="U", login="uu")
        self.client.post(reverse("dashboard:user_block", args=[user.pk]))
        self.client.logout()
        r = self.client.post(reverse("accounts:login"), {"username": "uu", "password": "Strong-pass-42"})
        self.assertContains(r, "заблокирован")
        self.client.force_login(self.admin)
        self.client.post(reverse("dashboard:user_block", args=[user.pk]))
        user.refresh_from_db()
        self.assertTrue(user.is_active)

    def test_delete_user(self):
        user = User.objects.create_user("u@b.kg", "Strong-pass-42", full_name="U")
        self.client.post(reverse("dashboard:user_delete", args=[user.pk]))
        self.assertFalse(User.objects.filter(pk=user.pk).exists())
        self.client.post(reverse("dashboard:user_delete", args=[self.admin.pk]))
        self.assertTrue(User.objects.filter(pk=self.admin.pk).exists())

    def test_grant_and_revoke_access(self):
        user = User.objects.create_user("u@b.kg", "Strong-pass-42", full_name="U")
        self.client.post(reverse("dashboard:user_access", args=[user.pk]), {"courses": [self.course.pk, self.other.pk]})
        self.assertEqual(user.enrollments.count(), 2)
        self.client.post(reverse("dashboard:user_access", args=[user.pk]), {"courses": [self.other.pk]})
        self.assertEqual(list(user.enrollments.values_list("course_id", flat=True)), [self.other.pk])

    def test_course_price(self):
        self.client.post(reverse("dashboard:course_create"), {"title": "Платный", "difficulty": "beginner", "price": "1990"})
        self.assertEqual(Course.objects.get(title="Платный").price, 1990)

    def test_library_upload_and_attach_to_lesson(self):
        r = self.client.post(
            reverse("dashboard:library_create"),
            {"title": "Вводное видео", "file": SimpleUploadedFile("intro.mp4", b"\x00\x00", content_type="video/mp4")},
        )
        self.assertRedirects(r, reverse("dashboard:library"))
        video = LibraryItem.objects.get()
        self.assertEqual(video.kind, "video")
        pdf = LibraryItem.objects.create(title="Конспект", file=SimpleUploadedFile("notes.pdf", b"%PDF"))
        self.assertEqual(pdf.kind, "file")
        lesson = self.course.levels.get().lessons.get()
        data = {
            "title": lesson.title, "order": 1, "description": "", "library_video": video.pk, "library_materials": [pdf.pk],
            "materials-TOTAL_FORMS": 0, "materials-INITIAL_FORMS": 0, "materials-MIN_NUM_FORMS": 0, "materials-MAX_NUM_FORMS": 1000,
        }
        self.client.post(reverse("dashboard:lesson_edit", args=[lesson.pk]), data)
        lesson.refresh_from_db()
        self.assertEqual(lesson.library_video, video)
        self.assertTrue(lesson.has_video)
        r = self.client.get(lesson.get_absolute_url())
        self.assertContains(r, video.file.url)
        self.assertContains(r, "Конспект")
        self.assertContains(self.client.get(reverse("dashboard:library")), "Вводное видео")

    def tearDown(self):
        import shutil

        shutil.rmtree("test_media", ignore_errors=True)


class ReportsTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user("adm@b.kg", "Strong-pass-42", full_name="Админ", is_staff=True)
        self.student = User.objects.create_user("s@b.kg", "Strong-pass-42", full_name="Айгерим Садыкова", login="aigerim")
        self.course = make_course(levels=1, lessons=1, questions=1)
        Enrollment.objects.create(user=self.student, course=self.course)
        self.client.force_login(self.admin)

    def test_report_page(self):
        r = self.client.get(reverse("dashboard:reports"), {"course": self.course.pk})
        self.assertContains(r, "Айгерим Садыкова")
        self.assertNotContains(r, "Админ</td>")

    def test_excel_report(self):
        from io import BytesIO

        from openpyxl import load_workbook

        r = self.client.get(reverse("dashboard:report_excel"))
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(BytesIO(r.content)).active
        values = [c.value for row in ws.iter_rows() for c in row]
        self.assertIn("Айгерим Садыкова", values)
        self.assertIn("aigerim", values)

    def test_pdf_report(self):
        r = self.client.get(reverse("dashboard:report_pdf"))
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_reports_forbidden_for_students(self):
        self.client.force_login(self.student)
        self.assertEqual(self.client.get(reverse("dashboard:report_excel")).status_code, 403)


class StudentTestsSectionTests(TestCase):
    def test_tests_page_lists_quizzes_of_accessible_courses(self):
        student = User.objects.create_user("s@b.kg", "Strong-pass-42", full_name="Студент")
        course = make_course(levels=2, lessons=1, questions=1)
        hidden = make_course(levels=1, lessons=1, questions=1)
        hidden.title = "Чужой курс"
        hidden.save()
        Enrollment.objects.create(user=student, course=course)
        self.client.force_login(student)
        r = self.client.get(reverse("quizzes:my_tests"))
        self.assertContains(r, "Уровень 1")
        self.assertContains(r, "досмотрите уроки")
        self.assertContains(r, "уровень закрыт")
        self.assertNotContains(r, "Чужой курс")
